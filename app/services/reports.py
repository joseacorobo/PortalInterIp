import io
import sqlite3
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

try:
    from database import get_db
except ImportError:
    from app.database import get_db

def _get_range_condition(range_filter: str):
    if range_filter == "today":
        return "DATE(tl.created_at) = DATE('now')"
    elif range_filter == "7days":
        return "DATE(tl.created_at) >= DATE('now', '-7 days')"
    elif range_filter == "month":
        return "STRFTIME('%Y-%m', tl.created_at) = STRFTIME('%Y-%m', 'now')"
    return "1=1"

def get_managerial_summary(area: str = "Todas", range_filter: str = "all") -> dict:
    conn = get_db()
    cur = conn.cursor()
    
    range_sql = _get_range_condition(range_filter)
    if area == "Todas":
        area_sql = "1=1"
        user_area_sql = "1=1"
        params = []
    elif area in ["Acceso", "Redes de Acceso"]:
        area_sql = "tl.area IN ('Soporte', 'Cabecera')"
        user_area_sql = "u.area IN ('Soporte', 'Cabecera')"
        params = []
    elif area in ["Servicios", "Servicios y Clientes"]:
        area_sql = "tl.area IN ('Telefonía')"
        user_area_sql = "u.area IN ('Telefonía')"
        params = []
    else:
        area_sql = "tl.area = ?"
        user_area_sql = "u.area = ?"
        params = [area]
    
    # 1. KPIs Globales
    cur.execute(f"""
    SELECT COUNT(tl.id) as total_tasks,
           COALESCE(SUM(tl.points), 0) as total_points,
           COALESCE(AVG(tl.net_duration), 0) as avg_mttr,
           COALESCE(SUM(tl.net_duration), 0) as total_net_min,
           COALESCE(SUM(tl.wait_minutes), 0) as total_wait_min
    FROM task_logs tl
    WHERE {range_sql} AND {area_sql}
    """, params)
    kpi_row = cur.fetchone()
    total_tasks = kpi_row[0] or 0
    total_points = kpi_row[1] or 0
    avg_mttr = round(kpi_row[2] or 0, 1)
    total_net_hours = round((kpi_row[3] or 0) / 60, 1)
    total_wait_hours = round((kpi_row[4] or 0) / 60, 1)
    
    # Cumplimiento SLA
    cur.execute(f"""
    SELECT COUNT(tl.id)
    FROM task_logs tl
    LEFT JOIN task_types tt ON tl.task_type_id = tt.id
    WHERE {range_sql} AND {area_sql} AND tl.net_duration <= COALESCE(tt.sla_minutes, 45)
    """, params)
    within_sla = cur.fetchone()[0] or 0
    sla_compliance = round((within_sla / total_tasks * 100), 1) if total_tasks > 0 else 100.0

    # 2. Desglose por Células / Áreas Funcionales
    areas_list = ["Soporte", "Cabecera", "Telefonía"]
    area_breakdown = []
    for a in areas_list:
        cur.execute(f"""
        SELECT COUNT(tl.id), COALESCE(SUM(tl.points), 0), COALESCE(AVG(tl.net_duration), 0)
        FROM task_logs tl
        WHERE {range_sql} AND tl.area = ?
        """, [a])
        r = cur.fetchone()
        a_tasks, a_points, a_mttr = r[0] or 0, r[1] or 0, round(r[2] or 0, 1)
        
        # Cantidad de especialistas
        cur.execute("SELECT COUNT(*) FROM users WHERE area = ?", [a])
        techs_count = cur.fetchone()[0] or 4
        pts_per_tech = round(a_points / techs_count, 1) if techs_count else 0
        
        if pts_per_tech <= 25:
            status = "Equilibrada"
            badge = "success"
        elif pts_per_tech <= 45:
            status = "Moderada"
            badge = "warning"
        else:
            status = "Saturada / Alerta"
            badge = "danger"
            
        area_breakdown.append({
            "area": a,
            "total_tasks": a_tasks,
            "total_points": a_points,
            "avg_mttr": a_mttr,
            "techs_count": techs_count,
            "pts_per_tech": pts_per_tech,
            "share_percent": round((a_points / total_points * 100), 1) if total_points > 0 else 0,
            "status": status,
            "badge": badge
        })

    # 3. Productividad Detallada por Especialista (Dinámico)
    cur.execute(f"""
    SELECT u.id, u.name, u.area, u.role, u.avatar,
           COUNT(tl.id) as tasks_count,
           COALESCE(SUM(tl.points), 0) as total_points,
           COALESCE(AVG(tl.net_duration), 0) as avg_mttr,
           SUM(CASE WHEN tt.code LIKE 'P1%' THEN 1 ELSE 0 END) as p1_count,
           SUM(CASE WHEN tt.code LIKE 'P2%' THEN 1 ELSE 0 END) as p2_count,
           SUM(CASE WHEN tt.code LIKE 'P3%' THEN 1 ELSE 0 END) as p3_count,
           SUM(CASE WHEN tt.code LIKE 'P4%' THEN 1 ELSE 0 END) as p4_count,
           SUM(CASE WHEN tt.code LIKE 'P5%' THEN 1 ELSE 0 END) as p5_count
    FROM users u
    LEFT JOIN task_logs tl ON u.id = tl.user_id AND {range_sql}
    LEFT JOIN task_types tt ON tl.task_type_id = tt.id
    WHERE {user_area_sql}
    GROUP BY u.id
    ORDER BY total_points DESC
    """, params)
    
    tech_rankings = []
    for row in cur.fetchall():
        pts = row[6] or 0
        if pts <= 25:
            st = "Equilibrada"
            st_color = "emerald"
        elif pts <= 45:
            st = "Moderada"
            st_color = "amber"
        else:
            st = "Alta"
            st_color = "red"
            
        tech_rankings.append({
            "id": row[0],
            "name": row[1],
            "area": row[2],
            "role": row[3],
            "avatar": row[4],
            "tasks_count": row[5] or 0,
            "total_points": pts,
            "avg_mttr": round(row[7] or 0, 1),
            "p1": row[8] or 0,
            "p2": row[9] or 0,
            "p3": row[10] or 0,
            "p4": row[11] or 0,
            "p5": row[12] or 0,
            "status": st,
            "status_color": st_color
        })

    # 4. Muestra de Registros de Auditoría
    cur.execute(f"""
    SELECT tl.ticket_code, tl.created_at, tl.duration_minutes, tl.wait_minutes, tl.net_duration,
           tl.points, tl.description, u.name as user_name, tl.area,
           et.subscriber_code, et.serial_pon, et.node_name, et.slot_pon, et.mac_address,
           tt.name as task_name, tt.code as task_code, tt.sla_minutes
    FROM task_logs tl
    LEFT JOIN users u ON tl.user_id = u.id
    LEFT JOIN task_types tt ON tl.task_type_id = tt.id
    LEFT JOIN email_tickets et ON tl.ticket_code = et.ticket_code
    WHERE {range_sql} AND {area_sql}
    ORDER BY tl.id DESC
    LIMIT 30
    """, params)
    
    audit_samples = []
    for r in cur.fetchall():
        sub = r[9] or "N/A"
        perm = f"P-{sub[:2]}" if sub != "N/A" and len(sub) >= 2 else "N/A"
        audit_samples.append({
            "ticket_code": r[0],
            "date": r[1],
            "duration": r[2],
            "wait_time": r[3],
            "net_duration": r[4],
            "points": r[5],
            "notes": r[6] or "",
            "user_name": r[7],
            "area": r[8],
            "subscriber_code": sub,
            "permisor": perm,
            "serial_pon": r[10] or "N/A",
            "node_name": r[11] or "N/A",
            "slot_pon": r[12] or "N/A",
            "mac_address": r[13] or "N/A",
            "task_name": r[14] or "Operación",
            "task_code": r[15] or "P2",
            "sla_minutes": r[16] or 45
        })

    conn.close()
    return {
        "kpis": {
            "total_points": total_points,
            "total_tasks": total_tasks,
            "avg_mttr": avg_mttr,
            "total_net_hours": total_net_hours,
            "total_wait_hours": total_wait_hours,
            "sla_compliance": sla_compliance
        },
        "area_breakdown": area_breakdown,
        "tech_rankings": tech_rankings,
        "audit_samples": audit_samples,
        "filters": {
            "area": area,
            "range": range_filter
        }
    }


def generate_excel_report(area: str = "Todas", range_filter: str = "all") -> io.BytesIO:
    """
    Genera un libro Excel con 3 pestañas profesionales para supervisión gerencial:
    1. Resumen Ejecutivo y Áreas
    2. Productividad por Especialista
    3. Log Detallado de Auditoría
    """
    data = get_managerial_summary(area, range_filter)
    wb = openpyxl.Workbook()
    
    # Fuentes y estilos corporativos SnowUI / Inter
    f_title = Font(name="Calibri", size=15, bold=True, color="1E3A8A")
    f_subtitle = Font(name="Calibri", size=10, italic=True, color="64748B")
    f_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    f_bold = Font(name="Calibri", size=11, bold=True, color="1E293B")
    f_regular = Font(name="Calibri", size=10, color="1E293B")
    f_kpi_val = Font(name="Calibri", size=18, bold=True, color="1E3A8A")
    f_kpi_lbl = Font(name="Calibri", size=9, bold=True, color="64748B")
    
    fill_navy = PatternFill("solid", fgColor="1E3A8A")
    fill_blue_head = PatternFill("solid", fgColor="2563EB")
    fill_ice = PatternFill("solid", fgColor="F8FAFC")
    fill_card = PatternFill("solid", fgColor="EFF6FF")
    
    border_thin = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # =========================================================================
    # HOJA 1: RESUMEN EJECUTIVO Y ÁREAS
    # =========================================================================
    ws1 = wb.active
    ws1.title = "Resumen Ejecutivo"
    ws1.views.sheetView[0].showGridLines = True
    
    # Encabezado Corporativo
    ws1.merge_cells("A1:G1")
    ws1["A1"] = "INTER TELECOMUNICACIONES — REPORTE GERENCIAL OPERACIONES IP"
    ws1["A1"].font = f_title
    ws1["A1"].alignment = align_left
    
    ws1.merge_cells("A2:G2")
    ws1["A2"] = f"División: Redes de Acceso y Aprovisionamiento | Filtro Área: {area} | Temporal: {range_filter.upper()} | Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws1["A2"].font = f_subtitle
    ws1["A2"].alignment = align_left
    
    # 4 Tarjetas KPI Superiores
    kpis = data["kpis"]
    kpi_cards = [
        ("PUNTOS TOTALES", f"{kpis['total_points']} pts", "A4:B5"),
        ("TICKETS RESUELTOS", f"{kpis['total_tasks']}", "C4:D5"),
        ("MTTR PROMEDIO NETO", f"{kpis['avg_mttr']} min", "E4:E5"),
        ("CUMPLIMIENTO SLA", f"{kpis['sla_compliance']}%", "F4:G5")
    ]
    for lbl, val, range_str in kpi_cards:
        cells = list(ws1[range_str])
        top_left = cells[0][0]
        top_left.value = f"{lbl}\n{val}"
        top_left.font = f_bold
        top_left.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        top_left.fill = fill_card
        for row in cells:
            for c in row:
                c.border = border_thin
                c.fill = fill_card

    # Tabla: Balance de Células Funcionales
    ws1["A7"] = "BALANCE DE CARGA POR CÉLULA FUNCIONAL"
    ws1["A7"].font = f_bold
    
    headers_ws1 = ["Célula / Área", "Especialistas", "Tickets Resueltos", "Puntos Totales", "% Participación", "MTTR Promedio (min)", "Estado de Carga"]
    for col_idx, h in enumerate(headers_ws1, 1):
        cell = ws1.cell(row=8, column=col_idx, value=h)
        cell.font = f_header
        cell.fill = fill_navy
        cell.alignment = align_center
        cell.border = border_thin
        
    curr_row = 9
    for ab in data["area_breakdown"]:
        ws1.cell(row=curr_row, column=1, value=ab["area"]).alignment = align_left
        ws1.cell(row=curr_row, column=2, value=ab["techs_count"]).alignment = align_center
        ws1.cell(row=curr_row, column=3, value=ab["total_tasks"]).alignment = align_right
        ws1.cell(row=curr_row, column=4, value=ab["total_points"]).alignment = align_right
        ws1.cell(row=curr_row, column=5, value=f"{ab['share_percent']}%").alignment = align_center
        ws1.cell(row=curr_row, column=6, value=ab["avg_mttr"]).alignment = align_right
        st_cell = ws1.cell(row=curr_row, column=7, value=ab["status"])
        st_cell.alignment = align_center
        if ab["status"] == "Equilibrada":
            st_cell.fill = PatternFill("solid", fgColor="D1FAE5")
        elif ab["status"] == "Moderada":
            st_cell.fill = PatternFill("solid", fgColor="FEF3C7")
        else:
            st_cell.fill = PatternFill("solid", fgColor="FEE2E2")
            
        for c in range(1, 8):
            ws1.cell(row=curr_row, column=c).border = border_thin
            ws1.cell(row=curr_row, column=c).font = f_regular
        curr_row += 1

    # =========================================================================
    # HOJA 2: PRODUCTIVIDAD POR ESPECIALISTA
    # =========================================================================
    ws2 = wb.create_sheet(title="Productividad Especialistas")
    ws2.views.sheetView[0].showGridLines = True
    
    ws2.merge_cells("A1:M1")
    ws2["A1"] = "RENDIMIENTO Y DESGLOSE POR ESPECIALISTA"
    ws2["A1"].font = f_title
    
    headers_ws2 = [
        "Especialista", "Célula / Área", "Rol Funcional", "Tickets", "Puntos Totales", 
        "Pts/Ticket", "P1 (1pt)", "P2 (2pts)", "P3 (3pts)", "P4 (5pts)", "P5 (8pts)", 
        "MTTR Prom (min)", "Estado de Carga"
    ]
    for col_idx, h in enumerate(headers_ws2, 1):
        cell = ws2.cell(row=3, column=col_idx, value=h)
        cell.font = f_header
        cell.fill = fill_blue_head
        cell.alignment = align_center
        cell.border = border_thin
        
    row_idx = 4
    for tech in data["tech_rankings"]:
        ws2.cell(row=row_idx, column=1, value=tech["name"]).alignment = align_left
        ws2.cell(row=row_idx, column=2, value=tech["area"]).alignment = align_center
        ws2.cell(row=row_idx, column=3, value=tech["role"]).alignment = align_left
        ws2.cell(row=row_idx, column=4, value=tech["tasks_count"]).alignment = align_right
        
        pts_cell = ws2.cell(row=row_idx, column=5, value=tech["total_points"])
        pts_cell.alignment = align_right
        pts_cell.font = f_bold
        
        pts_per_t = round(tech["total_points"] / tech["tasks_count"], 1) if tech["tasks_count"] > 0 else 0
        ws2.cell(row=row_idx, column=6, value=pts_per_t).alignment = align_right
        
        ws2.cell(row=row_idx, column=7, value=tech["p1"]).alignment = align_center
        ws2.cell(row=row_idx, column=8, value=tech["p2"]).alignment = align_center
        ws2.cell(row=row_idx, column=9, value=tech["p3"]).alignment = align_center
        ws2.cell(row=row_idx, column=10, value=tech["p4"]).alignment = align_center
        ws2.cell(row=row_idx, column=11, value=tech["p5"]).alignment = align_center
        
        ws2.cell(row=row_idx, column=12, value=tech["avg_mttr"]).alignment = align_right
        
        st_cell = ws2.cell(row=row_idx, column=13, value=tech["status"])
        st_cell.alignment = align_center
        if tech["status"] == "Equilibrada":
            st_cell.fill = PatternFill("solid", fgColor="D1FAE5")
        elif tech["status"] == "Moderada":
            st_cell.fill = PatternFill("solid", fgColor="FEF3C7")
        else:
            st_cell.fill = PatternFill("solid", fgColor="FEE2E2")
            
        for c in range(1, 14):
            ws2.cell(row=row_idx, column=c).border = border_thin
            if c != 5:
                ws2.cell(row=row_idx, column=c).font = f_regular
        row_idx += 1

    # =========================================================================
    # HOJA 3: LOG DETALLADO DE AUDITORÍA
    # =========================================================================
    ws3 = wb.create_sheet(title="Log Forense Auditoría")
    ws3.views.sheetView[0].showGridLines = True
    
    ws3.merge_cells("A1:R1")
    ws3["A1"] = "LOG DETALLADO DE AUDITORÍA Y TRAZABILIDAD TÉCNICA"
    ws3["A1"].font = f_title
    
    headers_ws3 = [
        "Ticket", "Fecha y Hora", "Duración Bruta (m)", "Pausa Terreno (m)", "Duración Neta (m)", 
        "Puntos", "Especialista", "Célula", "Abonado (10d)", "Permisor", "Serial PON (12c)", 
        "Nodo OLT", "Slot / PON", "MAC Abonado", "Tarea DERS", "Código Tarea", "SLA (min)", "Notas de Resolución"
    ]
    for col_idx, h in enumerate(headers_ws3, 1):
        cell = ws3.cell(row=3, column=col_idx, value=h)
        cell.font = f_header
        cell.fill = fill_navy
        cell.alignment = align_center
        cell.border = border_thin
        
    # Obtener todos los registros para auditoría completa
    conn = get_db()
    cur = conn.cursor()
    range_sql = _get_range_condition(range_filter)
    if area == "Todas":
        area_sql = "1=1"
        params = []
    elif area in ["Acceso", "Redes de Acceso"]:
        area_sql = "tl.area IN ('Soporte', 'Cabecera')"
        params = []
    elif area in ["Servicios", "Servicios y Clientes"]:
        area_sql = "tl.area IN ('Telefonía')"
        params = []
    else:
        area_sql = "tl.area = ?"
        params = [area]
    
    cur.execute(f"""
    SELECT tl.ticket_code, tl.created_at, tl.duration_minutes, tl.wait_minutes, tl.net_duration,
           tl.points, u.name as user_name, tl.area,
           et.subscriber_code, et.serial_pon, et.node_name, et.slot_pon, et.mac_address,
           tt.name as task_name, tt.code as task_code, tt.sla_minutes, tl.description
    FROM task_logs tl
    LEFT JOIN users u ON tl.user_id = u.id
    LEFT JOIN task_types tt ON tl.task_type_id = tt.id
    LEFT JOIN email_tickets et ON tl.ticket_code = et.ticket_code
    WHERE {range_sql} AND {area_sql}
    ORDER BY tl.id DESC
    """, params)
    
    audit_rows = cur.fetchall()
    conn.close()
    
    r_idx = 4
    for r in audit_rows:
        sub = r[8] or "N/A"
        perm = f"P-{sub[:2]}" if sub != "N/A" and len(sub) >= 2 else "N/A"
        
        ws3.cell(row=r_idx, column=1, value=r[0]).alignment = align_center
        ws3.cell(row=r_idx, column=2, value=r[1]).alignment = align_center
        ws3.cell(row=r_idx, column=3, value=r[2]).alignment = align_right
        ws3.cell(row=r_idx, column=4, value=r[3]).alignment = align_right
        ws3.cell(row=r_idx, column=5, value=r[4]).alignment = align_right
        ws3.cell(row=r_idx, column=6, value=r[5]).alignment = align_right
        ws3.cell(row=r_idx, column=7, value=r[6]).alignment = align_left
        ws3.cell(row=r_idx, column=8, value=r[7]).alignment = align_center
        ws3.cell(row=r_idx, column=9, value=sub).alignment = align_center
        ws3.cell(row=r_idx, column=10, value=perm).alignment = align_center
        ws3.cell(row=r_idx, column=11, value=r[9] or "N/A").alignment = align_center
        ws3.cell(row=r_idx, column=12, value=r[10] or "N/A").alignment = align_center
        ws3.cell(row=r_idx, column=13, value=r[11] or "N/A").alignment = align_left
        ws3.cell(row=r_idx, column=14, value=r[12] or "N/A").alignment = align_center
        ws3.cell(row=r_idx, column=15, value=r[13] or "N/A").alignment = align_left
        ws3.cell(row=r_idx, column=16, value=r[14] or "P2").alignment = align_center
        ws3.cell(row=r_idx, column=17, value=r[15] or 45).alignment = align_right
        ws3.cell(row=r_idx, column=18, value=r[16] or "").alignment = align_left
        
        for c in range(1, 19):
            ws3.cell(row=r_idx, column=c).border = border_thin
            ws3.cell(row=r_idx, column=c).font = f_regular
        r_idx += 1

    # Autoajustar ancho de columnas en todas las hojas
    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if "\n" in val:
                    val = max(val.split("\n"), key=len)
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def get_current_workload(area: str = "Todas") -> dict:
    """
    Calcula la carga de trabajo en tiempo real (Tickets EN PROGRESO)
    agrupados estructuradamente por Operador y por Área/Célula.
    Entrega un JSON limpio, tipado y optimizado para la supervisión operativa del NOC.
    """
    conn = get_db()
    conn.row_factory = sqlite3.Row
    cur = conn.cursor()
    
    try:
        # Filtros de área para tickets
        where_conditions = ["et.status = 'EN PROGRESO'"]
        params = []
        
        if area and area not in ["Todas", "Todas las Áreas", "Todas las Células"]:
            if area in ["Acceso", "Redes de Acceso"]:
                where_conditions.append("(et.area IN ('Soporte', 'Cabecera', 'Redes de Acceso y Aprovisionamiento') OR u.area IN ('Soporte', 'Cabecera', 'Acceso'))")
            elif area in ["Servicios", "Servicios y Clientes", "Telefonía"]:
                where_conditions.append("(et.area = 'Telefonía' OR u.area = 'Telefonía' OR d.codigo = 'TELEFONIA')")
            else:
                where_conditions.append("(et.area = ? OR u.area = ? OR d.nombre LIKE ?)")
                params.extend([area, area, f"%{area}%"])
                
        where_sql = " AND ".join(where_conditions)
        
        # 1. Obtener todos los tickets actualmente EN PROGRESO
        cur.execute(f"""
        SELECT et.id, et.ticket_code, et.subject, et.area as ticket_area,
               et.subscriber_code, et.serial_pon, et.node_name, et.slot_pon, et.mac_address,
               et.suggested_task_type_id, et.status, et.claimed_by_user_id, et.operador_id,
               COALESCE(et.fecha_inicio_atencion, et.claimed_at, et.created_at) as started_at,
               et.departamento_id,
               COALESCE(d.nombre, et.area) as departamento_nombre,
               COALESCE(d.codigo, 'ACCESO_APROV') as departamento_codigo,
               COALESCE(u.id, 0) as user_id,
               COALESCE(u.name, 'Sin Asignar') as user_name,
               COALESCE(u.avatar, 'OP') as user_avatar,
               COALESCE(u.role, 'ESPECIALISTA') as user_role,
               COALESCE(u.area, et.area) as user_area,
               tt.id as task_id, tt.code as task_code, tt.name as task_name,
               COALESCE(tt.points, 1) as task_points,
               COALESCE(tt.sla_minutes, 45) as sla_minutes
        FROM email_tickets et
        LEFT JOIN users u ON COALESCE(et.operador_id, et.claimed_by_user_id) = u.id
        LEFT JOIN departamentos d ON et.departamento_id = d.id
        LEFT JOIN task_types tt ON et.suggested_task_type_id = tt.id
        WHERE {where_sql}
        ORDER BY started_at ASC
        """, params)
        
        raw_tickets = [dict(r) for r in cur.fetchall()]
        
        now = datetime.now()
        processed_tickets = []
        
        for t in raw_tickets:
            started_str = t.get("started_at")
            elapsed_min = 1
            if started_str:
                for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M"):
                    try:
                        dt_started = datetime.strptime(started_str.split(".")[0], fmt)
                        diff_sec = max(0, (now - dt_started).total_seconds())
                        elapsed_min = max(1, round(diff_sec / 60))
                        break
                    except Exception:
                        continue
                    
            sla_min = t.get("sla_minutes") or 45
            sla_pct = min(100.0, round((elapsed_min / sla_min) * 100, 1))
            
            if sla_pct >= 100:
                sla_status = "breached"
                sla_label = "SLA Excedido"
                sla_color = "red"
            elif sla_pct >= 75:
                sla_status = "warning"
                sla_label = "SLA Crítico"
                sla_color = "amber"
            else:
                sla_status = "normal"
                sla_label = "En Tiempo"
                sla_color = "emerald"
                
            processed_tickets.append({
                "id": t["id"],
                "ticket_code": t["ticket_code"],
                "subject": t["subject"],
                "subscriber_code": t["subscriber_code"] or "N/A",
                "serial_pon": t["serial_pon"] or "N/A",
                "node_name": t["node_name"] or "N/A",
                "slot_pon": t["slot_pon"] or "N/A",
                "mac_address": t["mac_address"] or "N/A",
                "operator_id": t["user_id"],
                "operator_name": t["user_name"],
                "operator_avatar": t["user_avatar"],
                "operator_role": t["user_role"],
                "operator_area": t["user_area"],
                "department_id": t["departamento_id"],
                "department_name": t["departamento_nombre"],
                "department_code": t["departamento_codigo"],
                "task": {
                    "id": t["task_id"],
                    "code": t["task_code"] or "P2_STD",
                    "name": t["task_name"] or "Atención Estándar",
                    "points": t["task_points"],
                    "sla_minutes": sla_min
                },
                "points": t["task_points"],
                "started_at": started_str,
                "elapsed_minutes": elapsed_min,
                "sla_minutes": sla_min,
                "sla_percentage": sla_pct,
                "sla_status": sla_status,
                "sla_label": sla_label,
                "sla_color": sla_color
            })
            
        # 2. Agrupación por Operador
        user_area_where = []
        user_params = []
        if area and area not in ["Todas", "Todas las Áreas", "Todas las Células"]:
            if area in ["Acceso", "Redes de Acceso"]:
                user_area_where.append("u.area IN ('Soporte', 'Cabecera')")
            elif area in ["Servicios", "Servicios y Clientes"]:
                user_area_where.append("u.area IN ('Telefonía')")
            else:
                user_area_where.append("u.area = ?")
                user_params.append(area)
        user_where_sql = f"WHERE {' AND '.join(user_area_where)}" if user_area_where else ""
        
        cur.execute(f"""
        SELECT u.id, u.name, u.email, u.role, u.area, u.avatar, u.departamento_id,
               COALESCE(d.nombre, u.area) as departamento_nombre,
               COALESCE(d.codigo, 'ACCESO_APROV') as departamento_codigo
        FROM users u
        LEFT JOIN departamentos d ON u.departamento_id = d.id
        {user_where_sql}
        ORDER BY u.area ASC, u.name ASC
        """, user_params)
        all_operators = [dict(r) for r in cur.fetchall()]
        
        tickets_by_op = {}
        for t in processed_tickets:
            op_id = t["operator_id"]
            if op_id not in tickets_by_op:
                tickets_by_op[op_id] = []
            tickets_by_op[op_id].append(t)
            
        operators_workload = []
        for op in all_operators:
            op_id = op["id"]
            op_tickets = tickets_by_op.get(op_id, [])
            active_count = len(op_tickets)
            active_pts = sum(t["points"] for t in op_tickets)
            max_elapsed = max([t["elapsed_minutes"] for t in op_tickets], default=0)
            
            # Clasificación de nivel de saturación
            if active_count == 0:
                sat_level = "Disponible"
                sat_badge = "secondary"
                sat_color = "#10B981"
            elif active_count == 1 and active_pts <= 4:
                sat_level = "Equilibrada"
                sat_badge = "success"
                sat_color = "#1C58A8"
            elif active_count <= 2 or active_pts <= 8:
                sat_level = "Moderada"
                sat_badge = "warning"
                sat_color = "#F59E0B"
            else:
                sat_level = "Sobrecarga"
                sat_badge = "danger"
                sat_color = "#EF4444"
                
            operators_workload.append({
                "operator_id": op_id,
                "name": op["name"],
                "email": op["email"],
                "role": op["role"],
                "area": op["area"],
                "avatar": op["avatar"],
                "department_id": op["departamento_id"],
                "department_name": op["departamento_nombre"],
                "department_code": op["departamento_codigo"],
                "active_tickets_count": active_count,
                "active_points": active_pts,
                "max_elapsed_minutes": max_elapsed,
                "saturation_level": sat_level,
                "saturation_badge": sat_badge,
                "saturation_color": sat_color,
                "active_tickets": op_tickets
            })
            
        operators_workload.sort(key=lambda x: (x["active_tickets_count"] > 0, x["active_points"], x["active_tickets_count"]), reverse=True)
        
        # 3. Agrupación por Área / Célula
        areas_dict = {}
        for t in processed_tickets:
            a_name = t["operator_area"] or t["department_name"] or "General"
            if a_name not in areas_dict:
                areas_dict[a_name] = {
                    "area": a_name,
                    "department_id": t["department_id"],
                    "department_name": t["department_name"],
                    "department_code": t["department_code"],
                    "active_tickets_count": 0,
                    "active_points": 0,
                    "operators_set": set(),
                    "tickets": []
                }
            areas_dict[a_name]["active_tickets_count"] += 1
            areas_dict[a_name]["active_points"] += t["points"]
            if t["operator_id"]:
                areas_dict[a_name]["operators_set"].add(t["operator_id"])
            areas_dict[a_name]["tickets"].append(t)
            
        canonical_areas = ["Soporte", "Cabecera", "Telefonía"]
        for ca in canonical_areas:
            if ca not in areas_dict and (area == "Todas" or area == ca or (area in ["Acceso", "Redes de Acceso"] and ca in ["Soporte", "Cabecera"])):
                areas_dict[ca] = {
                    "area": ca,
                    "department_id": 1 if ca in ["Soporte", "Cabecera"] else 4,
                    "department_name": "Redes de Acceso y Aprovisionamiento" if ca in ["Soporte", "Cabecera"] else "Telefonía VoIP",
                    "department_code": "ACCESO_APROV" if ca in ["Soporte", "Cabecera"] else "TELEFONIA",
                    "active_tickets_count": 0,
                    "active_points": 0,
                    "operators_set": set(),
                    "tickets": []
                }
                
        total_active_tickets = len(processed_tickets)
        total_active_points = sum(t["points"] for t in processed_tickets)
        
        by_area = []
        for a_name, a_info in areas_dict.items():
            cnt = a_info["active_tickets_count"]
            pts = a_info["active_points"]
            share = round((pts / total_active_points * 100), 1) if total_active_points > 0 else 0
            by_area.append({
                "area": a_info["area"],
                "department_id": a_info["department_id"],
                "department_name": a_info["department_name"],
                "department_code": a_info["department_code"],
                "active_tickets_count": cnt,
                "active_points": pts,
                "active_operators_count": len(a_info["operators_set"]),
                "share_percentage": share,
                "saturation_status": "Alta" if pts > 15 else ("Moderada" if pts > 6 else "Normal"),
                "tickets": a_info["tickets"]
            })
        by_area.sort(key=lambda x: x["active_points"], reverse=True)
        
        # 4. Resumen Global
        active_ops_count = sum(1 for op in operators_workload if op["active_tickets_count"] > 0)
        avg_elapsed = round(sum(t["elapsed_minutes"] for t in processed_tickets) / total_active_tickets, 1) if total_active_tickets > 0 else 0
        bottleneck_area = by_area[0]["area"] if (by_area and by_area[0]["active_tickets_count"] > 0) else "Ninguna"
        most_loaded_op = operators_workload[0]["name"] if (operators_workload and operators_workload[0]["active_tickets_count"] > 0) else "Ninguno"
        
        summary = {
            "total_active_tickets": total_active_tickets,
            "total_active_points": total_active_points,
            "active_operators_count": active_ops_count,
            "total_available_operators": len(operators_workload),
            "avg_elapsed_minutes": avg_elapsed,
            "bottleneck_area": bottleneck_area,
            "most_loaded_operator": most_loaded_op,
            "system_status": "Alerta" if total_active_tickets > 10 else ("Atención" if total_active_tickets > 4 else "Óptimo")
        }
        
        return {
            "status": "success",
            "timestamp": now.isoformat(),
            "filter_area": area,
            "summary": summary,
            "by_operator": operators_workload,
            "by_area": by_area,
            "active_tickets_list": processed_tickets
        }
    finally:
        conn.close()
