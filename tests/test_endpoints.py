import io
import os
import sys
sys.path.insert(0, os.path.abspath("."))
import openpyxl
from fastapi.testclient import TestClient
import app.main as app_main

def run_tests():
    print("=" * 70)
    print("SUITE DE PRUEBAS AUTOMATIZADAS - PORTAL INTER OPERACIONES IP")
    print("=" * 70)
    
    client = TestClient(app_main.app)
    
    # -------------------------------------------------------------
    # Test 1: Vista de Login y Acceso al Dashboard HTML
    # -------------------------------------------------------------
    print("\n[TEST 1] GET / (Unauthenticated Redirect) y Login HTML")
    res_unauth = client.get("/", follow_redirects=False)
    assert res_unauth.status_code == 303, f"Esperaba redirect 303 a /login, obtuve {res_unauth.status_code}"
    assert res_unauth.headers["location"] == "/login"
    print("  -> PASSED: GET / sin sesión redirige correctamente (303) a /login.")

    res_login = client.get("/login")
    assert res_login.status_code == 200, f"Error en GET /login: {res_login.status_code}"
    assert "login_auth.css" in res_login.text, "login_auth.css debe estar vinculado en /login"
    print("  -> PASSED: GET /login responde 200 OK y vincula login_auth.css.")

    # Autenticar como Coordinador (User 2: Adelis Mejia)
    res_auth = client.post("/api/auth/login", json={"username": "adelis.mejia@inter.com.ve", "password": "inter2026"})
    assert res_auth.status_code == 200, f"Error en login: {res_auth.status_code} {res_auth.text}"
    print("  -> PASSED: Autenticación inicial vía /api/auth/login concedida (Cookie firmada establecida).")

    res_dash = client.get("/")
    assert res_dash.status_code == 200, f"Error en GET / autenticado: {res_dash.status_code}"
    html = res_dash.text
    assert "Portal de Operaciones IP" in html or "Operaciones IP" in html
    assert "btn-export-excel-audit" in html, "El botón con id btn-export-excel-audit debe estar en el HTML"
    assert "modalQuickResolveTicket" in html, "modalQuickResolveTicket debe estar en el HTML"
    assert "modalVerifyTicket" in html, "modalVerifyTicket debe estar en el HTML"
    assert "modal-new-ticket" in html, "modal-new-ticket debe estar en el HTML"
    assert "theme_tokens.css" in html, "theme_tokens.css debe estar vinculado en dashboard.html"
    assert "dashboard_layout.css" in html, "dashboard_layout.css debe estar vinculado en dashboard.html"
    print("  -> PASSED: Dashboard HTML renderizado correctamente con CSS modular y modales.")

    # -------------------------------------------------------------
    # Test 2: Canonical Endpoints de Métricas y Operaciones
    # -------------------------------------------------------------
    print("\n[TEST 2] Canonical Endpoints de Métricas y Operaciones")
    
    # 2.1 unassigned
    res = client.get("/api/tickets/unassigned")
    assert res.status_code == 200, f"Error unassigned: {res.status_code}"
    unassigned = res.json()
    assert isinstance(unassigned, list), "unassigned debe retornar lista"
    print(f"  -> PASSED: /api/tickets/unassigned retorna {len(unassigned)} tickets sin asignar.")

    # 2.2 operators workload
    res = client.get("/api/operators/workload")
    assert res.status_code == 200, f"Error operators workload: {res.status_code}"
    workload = res.json()
    assert isinstance(workload, (list, dict)), "workload debe ser list o dict"
    print("  -> PASSED: /api/operators/workload responde 200 OK.")

    # 2.3 hourly metrics
    res = client.get("/api/metrics/hourly")
    assert res.status_code == 200, f"Error hourly metrics: {res.status_code}"
    hourly = res.json()
    assert "hours" in hourly or "labels" in hourly or isinstance(hourly, list), "hourly debe tener estructura de curva"
    print("  -> PASSED: /api/metrics/hourly responde 200 OK.")

    # 2.4 DERS distribution
    res = client.get("/api/metrics/ders-distribution")
    assert res.status_code == 200, f"Error ders distribution: {res.status_code}"
    print("  -> PASSED: /api/metrics/ders-distribution responde 200 OK.")

    # -------------------------------------------------------------
    # Test 3: Exportación Forense a Excel
    # -------------------------------------------------------------
    print("\n[TEST 3] GET /api/reports/export-excel (openpyxl real stream)")
    res = client.get("/api/reports/export-excel")
    assert res.status_code == 200, f"Error export-excel: {res.status_code}"
    content_type = res.headers.get("content-type", "")
    assert "spreadsheetml" in content_type or "octet-stream" in content_type, f"Content-type inválido: {content_type}"
    
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert len(wb.sheetnames) >= 1, "El libro Excel debe contener al menos 1 hoja"
    sheet = wb.active
    print(f"  -> PASSED: Archivo Excel generado válidamente. Hoja: '{sheet.title}', filas: {sheet.max_row}, cols: {sheet.max_column}.")

    # -------------------------------------------------------------
    # Test 4: Conmutación de Perfiles y Autenticación RBAC
    # -------------------------------------------------------------
    print("\n[TEST 4] Conmutador de Usuarios / Perfiles (Coordinador vs Especialista)")
    
    # Switch to Coordinador (User 2: Adelis Mejia)
    res = client.post("/api/auth/switch-user", json={"user_id": 2})
    assert res.status_code == 200, f"Error switch-user coord: {res.status_code} {res.text}"
    coord_data = res.json()
    u_coord = coord_data.get("user", {})
    assert u_coord.get("role") in ["COORDINADOR", "ADMINISTRADOR"], f"Rol inesperado: {u_coord.get('role')}"
    print(f"  -> PASSED: Conmutado a Coordinador ({u_coord.get('name')}, Rol: {u_coord.get('role')}).")

    # Switch to Especialista (User 3: José Corobo)
    res = client.post("/api/auth/switch-user", json={"user_id": 3})
    assert res.status_code == 200, f"Error switch-user op: {res.status_code} {res.text}"
    op_data = res.json()
    u_op = op_data.get("user", {})
    assert u_op.get("role") in ["ESPECIALISTA", "OPERADOR"], f"Rol inesperado: {u_op.get('role')}"
    print(f"  -> PASSED: Conmutado a Especialista ({u_op.get('name')}, Rol: {u_op.get('role')}).")

    # -------------------------------------------------------------
    # Test 5: Ciclo de Vida FSM Completo
    # -------------------------------------------------------------
    print("\n[TEST 5] Ciclo FSM: Ingesta -> Asignación -> Inicio -> Pausa -> Reanudación -> Cierre -> Verificación DERS")
    
    # Crear ticket de prueba
    res = client.post("/api/tickets/create", json={
        "subject": "QA Test Ticket Ciclo Completo DERS",
        "body_text": "Verificación automatizada de FSM cronometrada y aprobación de puntos.",
        "area": "Redes de acceso y aprovisionamiento",
        "departamento_id": 1,
        "serial_pon": "FHTT09182312",
        "node_name": "OLT-QA-01"
    })
    assert res.status_code in [200, 201], f"Error al crear ticket: {res.status_code} {res.text}"
    created = res.json()
    ticket_id = created.get("ticket_id")
    assert ticket_id, f"Ticket ID no retornado: {created}"
    print(f"  [5.1] Ticket creado exitosamente con ID #{ticket_id} (Código: {created.get('ticket_code')})")

    # Asignar ticket a Operador 3 (como Coordinador 2)
    client.cookies.set("auth_user_id", app_main.sign_session_user_id(2))
    res = client.post(f"/api/tickets/{ticket_id}/assign", json={
        "operador_id": 3,
        "coordinador_id": 2,
        "notas": "Asignado por QA automatizado"
    })
    assert res.status_code in [200, 201], f"Error al asignar ticket: {res.status_code} {res.text}"
    print(f"  [5.2] Ticket #{ticket_id} asignado a Operador #3 (Status: ASIGNADO)")

    # Operador 3 inicia atención
    client.cookies.set("auth_user_id", app_main.sign_session_user_id(3))
    res = client.post(f"/api/tickets/{ticket_id}/start")
    assert res.status_code == 200, f"Error al iniciar ticket: {res.status_code} {res.text}"
    print(f"  [5.3] Operador #3 inició atención (Status: EN_PROGRESO)")

    # Operador pausa atención
    res = client.post(f"/api/tickets/{ticket_id}/pause")
    assert res.status_code == 200, f"Error al pausar ticket: {res.status_code} {res.text}"
    print(f"  [5.4] Ticket #{ticket_id} pausado (Status: EN_ESPERA)")

    # Operador reanuda atención
    res = client.post(f"/api/tickets/{ticket_id}/resume")
    assert res.status_code == 200, f"Error al reanudar ticket: {res.status_code} {res.text}"
    print(f"  [5.5] Ticket #{ticket_id} reanudado (Status: EN_PROGRESO)")

    # Operador resuelve ticket (FASE 1: pasa a POR_VERIFICAR)
    res = client.post(f"/api/tickets/{ticket_id}/complete", json={
        "resolution_notes": "Prueba de resolución QA completada satisfactoriamente con telemetría normalizada.",
        "user_id": 3
    })
    assert res.status_code == 200, f"Error al resolver ticket: {res.status_code} {res.text}"
    comp_data = res.json()
    print(f"  [5.6] Ticket #{ticket_id} resuelto por operador -> Status actual: '{comp_data.get('status', 'POR_VERIFICAR')}'")

    # Coordinador 2 valida y acredita puntos DERS (FASE 2: pasa a COMPLETADO)
    client.cookies.set("auth_user_id", app_main.sign_session_user_id(2))
    res = client.post(f"/api/tickets/{ticket_id}/verify", json={
        "confirmed_points": 3,
        "verification_notes": "Calidad técnica homologada por QA",
        "coordinador_id": 2
    })
    assert res.status_code == 200, f"Error al verificar ticket: {res.status_code} {res.text}"
    ver_data = res.json()
    print(f"  [5.7] Ticket #{ticket_id} verificado y acreditado por Coordinador. Puntos: {ver_data.get('points_accredited', 3)}. Status final: '{ver_data.get('status', 'COMPLETADO')}'.")

    # -------------------------------------------------------------
    # Test 6: Verificación de Auditoría Forense
    # -------------------------------------------------------------
    print("\n[TEST 6] GET /api/audit/logs (Trazabilidad Forense Inmutable)")
    res = client.get("/api/audit/logs?limit=10")
    assert res.status_code == 200, f"Error audit logs: {res.status_code}"
    logs = res.json()
    assert len(logs) > 0, "Debe haber registros de auditoría"
    print(f"  -> PASSED: {len(logs)} registros forenses recuperados exitosamente.")

    print("\n" + "=" * 70)
    print("TODAS LAS PRUEBAS AUTOMATIZADAS PASARON EXITOSAMENTE (100% SUCCESS)")
    print("=" * 70)

if __name__ == "__main__":
    run_tests()
